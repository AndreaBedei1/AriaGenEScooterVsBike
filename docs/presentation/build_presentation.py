"""Build docs/presentation/AriaGenEScooter_overview.pptx (8 slides, Italian).

Short overview deck for the project lead. Requires the 'presentation' extra
(python-pptx); run from the repository root:

    python docs/presentation/build_presentation.py

All diagrams are native, editable PowerPoint shapes. The two images come from
the SYNTHETIC session (semantic-gaze illustration rendered here) and from the
MVP diagnostic figure (docs/figures/synthetic_overview.png); both are labelled
as synthetic.
"""

from __future__ import annotations

import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT / "src"))
from ariagen_escooter.segmentation import LabelMapStore  # noqa: E402

OUT = HERE / "AriaGenEScooter_overview.pptx"
ASSETS = HERE / "assets"
SYN = ROOT / "data" / "synthetic" / "SYN001_S01" / "SYN001_S01_R01"
T0_NS = 12_345_000_000_000

# Palette: asphalt + cycle-path teal + "gaze" amber
ASPHALT = RGBColor(0x1F, 0x2B, 0x38)
TEAL = RGBColor(0x1A, 0x9E, 0x8F)
TEAL_DARK = RGBColor(0x11, 0x6D, 0x63)
AMBER = RGBColor(0xF2, 0xA3, 0x3A)
CARD = RGBColor(0xEC, 0xF5, 0xF3)
CARD_DARK = RGBColor(0x2C, 0x3B, 0x4B)
TEXT = RGBColor(0x1F, 0x2B, 0x38)
MUTED = RGBColor(0x55, 0x62, 0x70)
SOFT = RGBColor(0xC6, 0xD5, 0xD2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROAD = RGBColor(0x7A, 0x84, 0x8F)
LANE = RGBColor(0xC8, 0x4B, 0x3F)
WALK = RGBColor(0xD9, 0xD0, 0xC1)
GREEN = RGBColor(0x7F, 0xB0, 0x69)
PED = RGBColor(0x2E, 0x6F, 0xD8)
FONT = "Calibri"


# --------------------------------------------------------------------------- #
# Low-level helpers
# --------------------------------------------------------------------------- #
def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.5, name=None, rot=0.0, dash=None):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
        if dash is not None:
            shp.line.dash_style = dash
    shp.shadow.inherit = False
    if rot:
        shp.rotation = rot
    if name:
        shp.name = name
    return shp


def text(
    slide,
    x,
    y,
    w,
    h,
    paragraphs,
    size=16,
    color=TEXT,
    bold=False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    italic=False,
    space_after=0,
    name=None,
):
    """Text box without padding. ``paragraphs``: str | list[str | list[(text, opts)]]."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paragraphs, str):
        paragraphs = [paragraphs]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = [(para, {})] if isinstance(para, str) else para
        for run_text, opts in runs:
            r = p.add_run()
            r.text = run_text
            f = r.font
            f.name = FONT
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", italic)
            f.color.rgb = opts.get("color", color)
    if name:
        tb.name = name
    return tb


def line(slide, x1, y1, x2, y2, color, width=2.0, arrow=False, dash=None, name=None):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    if dash is not None:
        conn.line.dash_style = dash
    if arrow:
        ln = conn.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    if name:
        conn.name = name
    return conn


def polyline(slide, points, color, width=3.0, name=None):
    pts = [(Inches(px), Inches(py)) for px, py in points]
    builder = slide.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
    builder.add_line_segments(pts[1:], close=False)
    shp = builder.convert_to_shape()
    shp.fill.background()
    shp.line.color.rgb = color
    shp.line.width = Pt(width)
    shp.shadow.inherit = False
    if name:
        shp.name = name
    return shp


def background(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def title(slide, label, dark=False):
    text(slide, 0.6, 0.45, 12.1, 0.8, label, size=36, bold=True, color=WHITE if dark else TEXT, name="Title")


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


def badge(slide, cx, cy, d, color=TEAL, name=None):
    return shape(slide, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill=color, name=name)


# --------------------------------------------------------------------------- #
# Pictograms (white on a coloured circle), all native shapes
# --------------------------------------------------------------------------- #
def icon_glasses(s, cx, cy, k=1.0):
    for dx in (-0.21, 0.21):
        shape(s, MSO_SHAPE.OVAL, cx + dx * k - 0.16 * k, cy - 0.13 * k, 0.32 * k, 0.26 * k, line=WHITE, line_w=2.5)
    line(s, cx - 0.05 * k, cy - 0.04 * k, cx + 0.05 * k, cy - 0.04 * k, WHITE, 2.5)
    line(s, cx - 0.37 * k, cy - 0.06 * k, cx - 0.47 * k, cy - 0.12 * k, WHITE, 2.5)
    line(s, cx + 0.37 * k, cy - 0.06 * k, cx + 0.47 * k, cy - 0.12 * k, WHITE, 2.5)


def icon_eye(s, cx, cy, k=1.0):
    shape(s, MSO_SHAPE.OVAL, cx - 0.4 * k, cy - 0.21 * k, 0.8 * k, 0.42 * k, fill=WHITE)
    shape(s, MSO_SHAPE.OVAL, cx - 0.16 * k, cy - 0.16 * k, 0.32 * k, 0.32 * k, fill=AMBER)
    shape(s, MSO_SHAPE.OVAL, cx - 0.07 * k, cy - 0.07 * k, 0.14 * k, 0.14 * k, fill=ASPHALT)


def icon_video(s, cx, cy, k=1.0):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.33 * k, cy - 0.23 * k, 0.66 * k, 0.46 * k, fill=WHITE)
    shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, cx - 0.1 * k, cy - 0.11 * k, 0.22 * k, 0.22 * k, fill=TEAL, rot=90)


def icon_trajectory(s, cx, cy, k=1.0):
    xs = np.linspace(-0.36, 0.3, 12)
    pts = [(cx + x * k, cy + 0.14 * k * np.sin(x * 9.0)) for x in xs]
    polyline(s, pts, WHITE, 3.0)
    end = pts[-1]
    shape(s, MSO_SHAPE.OVAL, end[0] - 0.07 * k, end[1] - 0.07 * k, 0.14 * k, 0.14 * k, fill=AMBER)


def icon_pin(s, cx, cy, k=1.0):
    shape(s, MSO_SHAPE.TEAR, cx - 0.22 * k, cy - 0.3 * k, 0.44 * k, 0.44 * k, fill=WHITE, rot=135)
    shape(s, MSO_SHAPE.OVAL, cx - 0.08 * k, cy - 0.16 * k, 0.16 * k, 0.16 * k, fill=TEAL)


def icon_layers(s, cx, cy, k=1.0):
    for i, col in enumerate((SOFT, WHITE, WHITE)):
        shape(s, MSO_SHAPE.PARALLELOGRAM, cx - 0.34 * k, cy + (0.12 - 0.13 * i) * k, 0.68 * k, 0.16 * k, fill=col)


def icon_speed(s, cx, cy, k=1.0):
    shape(s, MSO_SHAPE.BLOCK_ARC, cx - 0.34 * k, cy - 0.3 * k, 0.68 * k, 0.68 * k, fill=WHITE)
    line(s, cx, cy + 0.04 * k, cx + 0.2 * k, cy - 0.16 * k, AMBER, 3.0)


def icon_interaction(s, cx, cy, k=1.0):
    shape(s, MSO_SHAPE.OVAL, cx - 0.4 * k, cy - 0.11 * k, 0.22 * k, 0.22 * k, fill=WHITE)
    shape(s, MSO_SHAPE.OVAL, cx + 0.18 * k, cy - 0.11 * k, 0.22 * k, 0.22 * k, fill=AMBER)
    shape(s, MSO_SHAPE.LEFT_RIGHT_ARROW, cx - 0.16 * k, cy - 0.07 * k, 0.32 * k, 0.14 * k, fill=WHITE)


def icon_no(s, cx, cy, k=1.0):
    shape(s, MSO_SHAPE.NO_SYMBOL, cx - 0.3 * k, cy - 0.3 * k, 0.6 * k, 0.6 * k, fill=WHITE)


def icon_route(s, cx, cy, k=1.0):
    pts = [
        (cx - 0.32 * k, cy + 0.22 * k),
        (cx - 0.1 * k, cy + 0.22 * k),
        (cx - 0.1 * k, cy - 0.05 * k),
        (cx + 0.14 * k, cy - 0.05 * k),
        (cx + 0.14 * k, cy - 0.22 * k),
        (cx + 0.32 * k, cy - 0.22 * k),
    ]
    polyline(s, pts, WHITE, 3.0)
    for px, py in (pts[0], pts[-1]):
        shape(s, MSO_SHAPE.OVAL, px - 0.07 * k, py - 0.07 * k, 0.14 * k, 0.14 * k, fill=AMBER)


def icon_blur(s, cx, cy, k=1.0):
    shades = [WHITE, SOFT, WHITE, SOFT, WHITE, SOFT, WHITE, SOFT, WHITE]
    for i in range(3):
        for j in range(3):
            shape(
                s,
                MSO_SHAPE.RECTANGLE,
                cx + (-0.27 + 0.19 * j) * k,
                cy + (-0.27 + 0.19 * i) * k,
                0.16 * k,
                0.16 * k,
                fill=shades[i * 3 + j],
            )


def icon_id(s, cx, cy, k=1.0):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.36 * k, cy - 0.21 * k, 0.72 * k, 0.42 * k, fill=WHITE)
    text(
        s,
        cx - 0.36 * k,
        cy - 0.21 * k,
        0.72 * k,
        0.42 * k,
        "P001",
        size=14,
        bold=True,
        color=TEAL_DARK,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )


# --------------------------------------------------------------------------- #
# Infrastructure pictograms (top view, 1.5 x 1.5 in)
# --------------------------------------------------------------------------- #
def pict_frame(s, x, y, size, color):
    shape(s, MSO_SHAPE.RECTANGLE, x, y, size, size, fill=color)


def scooter_dot(s, x, y):
    shape(s, MSO_SHAPE.OVAL, x - 0.09, y - 0.09, 0.18, 0.18, fill=TEAL, line=WHITE, line_w=1.5)


def pict_protected(s, x, y, a=1.5):
    widths = [(0.4, ROAD), (0.06, ASPHALT), (0.3, LANE), (0.06, ASPHALT), (0.18, WALK)]
    cursor = x
    for frac, col in widths:
        shape(s, MSO_SHAPE.RECTANGLE, cursor, y, frac * a, a, fill=col)
        cursor += frac * a
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.1 * a, y + 0.15 * a, 0.2 * a, 0.32 * a, fill=ASPHALT)
    scooter_dot(s, x + 0.61 * a, y + 0.62 * a)


def pict_shared(s, x, y, a=1.5):
    pict_frame(s, x, y, a, GREEN)
    shape(s, MSO_SHAPE.RECTANGLE, x + 0.2 * a, y, 0.6 * a, a, fill=WALK)
    shape(s, MSO_SHAPE.OVAL, x + 0.32 * a - 0.08, y + 0.3 * a - 0.08, 0.16, 0.16, fill=PED)
    shape(s, MSO_SHAPE.OVAL, x + 0.62 * a - 0.08, y + 0.45 * a - 0.08, 0.16, 0.16, fill=PED)
    scooter_dot(s, x + 0.45 * a, y + 0.75 * a)


def pict_road(s, x, y, a=1.5):
    pict_frame(s, x, y, a, ROAD)
    for i in range(4):
        shape(s, MSO_SHAPE.RECTANGLE, x + 0.48 * a, y + (0.05 + 0.26 * i) * a, 0.04 * a, 0.14 * a, fill=WHITE)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.62 * a, y + 0.18 * a, 0.2 * a, 0.32 * a, fill=ASPHALT)
    scooter_dot(s, x + 0.75 * a, y + 0.72 * a)


def pict_crossing(s, x, y, a=1.5):
    pict_frame(s, x, y, a, WALK)
    shape(s, MSO_SHAPE.RECTANGLE, x, y + 0.25 * a, a, 0.5 * a, fill=ROAD)
    for i in range(6):
        shape(s, MSO_SHAPE.RECTANGLE, x + (0.1 + 0.14 * i) * a, y + 0.3 * a, 0.08 * a, 0.4 * a, fill=WHITE)
    scooter_dot(s, x + 0.5 * a, y + 0.88 * a)


def pict_intersection(s, x, y, a=1.5):
    pict_frame(s, x, y, a, WALK)
    shape(s, MSO_SHAPE.RECTANGLE, x + 0.32 * a, y, 0.36 * a, a, fill=ROAD)
    shape(s, MSO_SHAPE.RECTANGLE, x, y + 0.32 * a, a, 0.36 * a, fill=ROAD)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.06 * a, y + 0.36 * a, 0.22 * a, 0.12 * a, fill=ASPHALT)
    scooter_dot(s, x + 0.58 * a, y + 0.82 * a)


def pict_roundabout(s, x, y, a=1.5):
    pict_frame(s, x, y, a, WALK)
    shape(s, MSO_SHAPE.DONUT, x + 0.12 * a, y + 0.12 * a, 0.76 * a, 0.76 * a, fill=ROAD)
    shape(s, MSO_SHAPE.OVAL, x + 0.33 * a, y + 0.33 * a, 0.34 * a, 0.34 * a, fill=GREEN)
    scooter_dot(s, x + 0.5 * a, y + 0.82 * a)


# --------------------------------------------------------------------------- #
# Image assets
# --------------------------------------------------------------------------- #
CLASS_COLORS = {
    "road": "#7A848F",
    "cycle_lane": "#C84B3F",
    "sidewalk": "#D9D0C1",
    "crossing": "#F4F4F4",
    "curb": "#4A4F55",
    "pedestrian": "#2E6FD8",
    "cyclist": "#5DB7DE",
    "car": "#E07A2E",
    "parked_vehicle": "#B9A44C",
    "traffic_sign": "#F2C744",
    "traffic_light": "#3E8E41",
    "vegetation": "#8CBF6E",
    "building": "#B8B2CC",
    "sky": "#CFE3F2",
    "unknown": "#FFFFFF",
    "other": "#DDDDDD",
}
LABELS_IT = {
    "pedestrian": "pedone",
    "cycle_lane": "pista ciclabile",
    "road": "strada",
    "sidewalk": "marciapiede",
    "building": "edifici",
    "traffic_sign": "segnale",
}


def render_scene_image(path: Path, t_rel_s: float = 46.5) -> Path:
    store = LabelMapStore.load(SYN / "semantic" / "labels.npz")
    frames = pd.read_csv(SYN / "streams" / "frames.csv")
    gaze = pd.read_csv(SYN / "streams" / "gaze.csv")
    k = int(np.argmin(np.abs(frames["t_ns"].to_numpy() - (T0_NS + t_rel_s * 1e9))))
    frame_t = int(frames["t_ns"].iat[k])
    labels = store.labels[store.row(int(frames["frame_index"].iat[k]))]
    names = store.class_names
    lut = np.array([matplotlib.colors.to_rgb(CLASS_COLORS.get(n, "#DDDDDD")) for n in names])
    width, height = store.source_size
    near = gaze[(np.abs(gaze["t_ns"] - frame_t) <= 50_000_000) & gaze["valid"]]
    u, v = float(near["u_px"].mean()), float(near["v_px"].mean())

    fig = plt.figure(figsize=(8, 6), dpi=150)
    ax = fig.add_axes((0, 0, 1, 1))
    ax.imshow(lut[labels], extent=(-0.5, width - 0.5, height - 0.5, -0.5), interpolation="nearest")
    ax.add_patch(plt.Circle((u, v), 95, fill=False, lw=6, color="#F2A33A"))
    ax.add_patch(plt.Circle((u, v), 22, color="#F2A33A"))
    map_h, map_w = labels.shape
    for cls, label in LABELS_IT.items():
        cid = names.index(cls)
        rows, cols = np.nonzero(labels == cid)
        if rows.size < 4:
            continue
        r, c = np.median(rows), np.median(cols)
        x, y = (c + 0.5) * width / map_w, (r + 0.5) * height / map_h
        if cls == "pedestrian":
            x, y = x + 330, y - 230
        ax.text(
            x,
            y,
            label,
            fontsize=17,
            ha="center",
            va="center",
            color="#1F2B38",
            bbox={"boxstyle": "round,pad=0.3", "fc": "white", "ec": "none", "alpha": 0.9},
        )
    ax.text(
        width * 0.02,
        height * 0.96,
        "DATI SINTETICI",
        fontsize=15,
        color="#1F2B38",
        fontweight="bold",
        bbox={"boxstyle": "round,pad=0.3", "fc": "white", "ec": "none", "alpha": 0.9},
    )
    ax.set_xlim(-0.5, width - 0.5)
    ax.set_ylim(height - 0.5, -0.5)
    ax.axis("off")
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


# --------------------------------------------------------------------------- #
# Slides
# --------------------------------------------------------------------------- #
def slide_idea(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, ASPHALT)
    text(s, 0.8, 1.55, 7.6, 0.4, "AriaGenEScooter  |  IDEA", size=16, bold=True, color=TEAL)
    # explicit lines: never break "e-scooter" at its hyphen
    text(
        s,
        0.8,
        2.1,
        7.6,
        2.2,
        ["Studiare guida e attenzione", "visiva durante l'uso reale", "dell'e-scooter"],
        size=40,
        bold=True,
        color=WHITE,
        name="Main idea",
    )
    text(
        s,
        0.8,
        4.45,
        7.7,
        0.9,
        "Studio naturalistico in città con gli occhiali di ricerca Meta Project Aria Gen 2",
        size=20,
        color=SOFT,
    )
    text(s, 0.8, 6.55, 7.6, 0.4, "Bozza per discussione interna  |  ottobre 2026", size=12, color=SOFT)
    # gaze-target motif
    cx, cy = 10.55, 3.75
    for d, col in ((3.6, CARD_DARK), (2.6, TEAL_DARK), (1.6, TEAL)):
        shape(s, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill=col)
    shape(s, MSO_SHAPE.OVAL, cx - 0.36, cy - 0.36, 0.72, 0.72, fill=AMBER, name="Gaze point")
    notes(
        s,
        "Idea in una frase: osservare nel traffico reale, senza provocare nulla, dove guarda chi guida un "
        "monopattino e come guida, usando gli occhiali di ricerca Aria Gen 2. Non è uno studio su incidenti.",
    )


def slide_recording(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Cosa registriamo")
    cards = [
        ("Aria Gen 2", "occhiali di ricerca con sensori", icon_glasses),
        ("Sguardo", "dove guarda il rider (30 volte al secondo)", icon_eye),
        ("Video RGB", "la scena davanti al rider", icon_video),
        ("IMU / VIO", "movimento e traiettoria", icon_trajectory),
        ("Posizione", "GPS, quando disponibile", icon_pin),
    ]
    w, gap, x0, y0 = 2.2, 0.27, 0.6, 1.75
    for i, (label, caption, icon) in enumerate(cards):
        x = x0 + i * (w + gap)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y0, w, 3.75, fill=CARD, name=f"Card {label}")
        badge(s, x + w / 2, y0 + 1.0, 1.3)
        icon(s, x + w / 2, y0 + 1.0, 1.1)
        text(s, x + 0.2, y0 + 2.0, w - 0.4, 0.5, label, size=20, bold=True, align=PP_ALIGN.CENTER)
        text(s, x + 0.2, y0 + 2.55, w - 0.4, 1.0, caption, size=15, color=MUTED, align=PP_ALIGN.CENTER)
    text(
        s,
        0.6,
        6.0,
        12.1,
        0.6,
        [
            [
                ("Un solo orologio per tutti i sensori: ", {"bold": True}),
                ("sguardo, video e movimento sono sincronizzati con precisione.", {}),
            ]
        ],
        size=16,
        color=TEXT,
    )
    notes(
        s,
        "Gli occhiali registrano video della scena, direzione dello sguardo, movimento della testa e traiettoria; "
        "il GPS quando disponibile. Audio e immagini degli occhi non servono all'analisi.",
    )


def slide_where(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Dove: contesti urbani diversi")
    tiles = [
        ("Pista ciclabile protetta", "protected cycle path", pict_protected),
        ("Percorso condiviso", "shared pedestrian-cycle path", pict_shared),
        ("Strada con traffico", "road shared with traffic", pict_road),
        ("Attraversamenti", "crossings", pict_crossing),
        ("Intersezioni", "intersections", pict_intersection),
        ("Rotatorie", "roundabouts", pict_roundabout),
    ]
    tw, th, gx, gy, x0, y0 = 3.85, 1.95, 0.27, 0.3, 0.6, 1.6
    for i, (label, en, pict) in enumerate(tiles):
        x = x0 + (i % 3) * (tw + gx)
        y = y0 + (i // 3) * (th + gy)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, tw, th, fill=CARD, name=f"Tile {en}")
        pict(s, x + 0.22, y + 0.22, 1.5)
        text(s, x + 1.95, y + 0.5, tw - 2.1, 0.75, label, size=18, bold=True)
        text(s, x + 1.95, y + 1.25, tw - 2.1, 0.45, en, size=13, color=MUTED, italic=True)
    text(
        s,
        0.6,
        6.45,
        12.1,
        0.6,
        [
            [
                ("Unità di analisi: ", {"bold": True}),
                (
                    "il tratto di percorso omogeneo; pedoni, ciclisti e veicoli sono interazioni che vi si sovrappongono.",
                    {},
                ),
            ]
        ],
        size=15,
    )
    notes(
        s,
        "Le rotatorie sono un contesto importante ma non il centro dello studio: la maggior parte dei dati "
        "verrà da piste ciclabili e percorsi urbani. Il percorso definitivo deve rispettare le norme sui "
        "monopattini (da verificare).",
    )


def slide_analysis(prs, scene_png: Path):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Cosa analizziamo")
    s.shapes.add_picture(
        str(scene_png), Inches(0.6), Inches(1.6), Inches(5.9), Inches(4.425)
    ).name = "Semantic gaze example"
    text(
        s,
        0.6,
        6.15,
        5.9,
        0.6,
        "Esempio sintetico: il cerchio arancione è lo sguardo, che cade su un pedone",
        size=13,
        color=MUTED,
        italic=True,
    )
    rows = [
        ("Scena semantica", "strada, pista ciclabile, pedoni, veicoli, segnaletica", icon_layers),
        ("Sguardo (semantic gaze)", "cosa osserva il rider e per quanto tempo", icon_eye),
        ("Dinamica del rider", "velocità, rallentamenti, stabilità della traiettoria", icon_speed),
        ("Interazioni", "pedoni, ciclisti, altri monopattini, veicoli, ostacoli", icon_interaction),
    ]
    for i, (label, caption, icon) in enumerate(rows):
        y = 1.75 + i * 1.12
        badge(s, 7.35, y + 0.42, 0.85)
        icon(s, 7.35, y + 0.42, 0.72)
        text(s, 8.05, y + 0.02, 4.7, 0.45, label, size=20, bold=True)
        text(s, 8.05, y + 0.48, 4.7, 0.5, caption, size=15, color=MUTED)
    notes(
        s,
        "Ogni campione di sguardo viene associato all'oggetto o alla superficie che il rider sta guardando "
        "(semantic gaze). Da qui: tempi di osservazione per categoria, sguardi laterali, sguardo in avanti.",
    )


def slide_perception_action(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Perception → Action")
    steps = [
        ("VEDERE", "see", AMBER, ASPHALT),
        ("CAPIRE / INTERAGIRE", "understand / interact", TEAL, WHITE),
        ("REAGIRE", "react", TEAL_DARK, WHITE),
    ]
    x, w, gap = 0.6, 3.95, 0.12
    for label, en, fill, fg in steps:
        shape(s, MSO_SHAPE.CHEVRON, x, 1.75, w, 1.35, fill=fill, name=f"Step {en}")
        text(s, x + 0.55, 1.95, w - 1.1, 0.55, label, size=22, bold=True, color=fg, align=PP_ALIGN.CENTER)
        text(s, x + 0.55, 2.5, w - 1.1, 0.4, en, size=14, italic=True, color=fg, align=PP_ALIGN.CENTER)
        x += w + gap
    # timeline (synthetic example: pedestrian crossing the cycle path)
    y = 4.65
    line(s, 0.9, y, 12.4, y, ASPHALT, 3.0, arrow=True, name="Timeline")
    gaze_x, onset_x, peak_x, pass_x = 1.95, 4.55, 7.05, 10.6
    marks = [
        (gaze_x, "primo sguardo al pedone", "43,6 s", AMBER),
        (onset_x, "inizio rallentamento", "44,1 s", TEAL),
        (peak_x, "decelerazione massima", "", TEAL_DARK),
        (pass_x, "passaggio vicino al pedone", "49,0 s", ASPHALT),
    ]
    for mx, label, value, col in marks:
        shape(s, MSO_SHAPE.OVAL, mx - 0.16, y - 0.16, 0.32, 0.32, fill=col)
        text(
            s,
            mx - 1.15,
            y + 0.3,
            2.3,
            0.75,
            [[(label, {"bold": True})]] + ([value] if value else []),
            size=14,
            align=PP_ALIGN.CENTER,
        )
    # spans: latency bracket just above the first two markers
    by = 4.2
    line(s, gaze_x, by, onset_x, by, AMBER, 2.5)
    for bx in (gaze_x, onset_x):
        line(s, bx, by - 0.1, bx, by + 0.22, AMBER, 2.5)
    text(
        s,
        gaze_x - 0.6,
        by - 0.48,
        onset_x - gaze_x + 1.2,
        0.4,
        "latenza percezione → azione: 0,5 s",
        size=14,
        bold=True,
        color=TEXT,
        align=PP_ALIGN.CENTER,
    )
    line(s, gaze_x, 6.2, pass_x, 6.2, MUTED, 1.5, dash=MSO_LINE_DASH_STYLE.DASH)
    text(s, 4.0, 6.3, 4.6, 0.4, "anticipo visivo: 5,4 s", size=14, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    text(
        s,
        0.6,
        6.85,
        12.1,
        0.35,
        'Esempio su dati sintetici. Termini neutri: interazione, evento (non "pericolo").',
        size=12,
        color=MUTED,
        italic=True,
    )
    notes(
        s,
        "Domanda centrale: quanto prima il rider vede un pedone o un veicolo, e quanto dopo reagisce. "
        "Misuriamo il primo sguardo sul bersaglio, l'inizio della decelerazione, la decelerazione massima "
        "e il momento di massimo avvicinamento.",
    )


def slide_comparisons(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Confronti")
    cols = ["pista protetta", "condiviso", "strada", "rotatoria"]
    riders = ["Rider A", "Rider B", "Rider C", "Rider D"]
    gx0, gy0, dx, dy = 2.0, 2.2, 1.15, 0.95
    for j, c in enumerate(cols):
        text(s, gx0 + j * dx - 0.55, gy0 - 0.65, 1.1, 0.4, c, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    for i, r in enumerate(riders):
        text(s, 0.6, gy0 + i * dy - 0.18, 1.0, 0.4, r, size=14, bold=True)
        for j in range(len(cols)):
            col = AMBER if i == 1 else (TEAL if j == 2 else SOFT)
            shape(s, MSO_SHAPE.OVAL, gx0 + j * dx - 0.2, gy0 + i * dy - 0.2, 0.4, 0.4, fill=col)
    shape(
        s,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        gx0 - 0.45,
        gy0 + dy - 0.38,
        3 * dx + 0.9,
        0.76,
        line=AMBER,
        line_w=2.5,
        name="Same rider row",
    )
    shape(
        s,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        gx0 + 2 * dx - 0.38,
        gy0 - 0.38,
        0.76,
        3 * dy + 0.76,
        line=TEAL,
        line_w=2.5,
        name="Same infrastructure column",
    )
    items = [
        (AMBER, "Stesso rider, infrastrutture diverse", "come cambia lo stesso conducente da un contesto all'altro"),
        (TEAL, "Rider diversi, stessa infrastruttura", "cosa è comune a tutti in uno stesso contesto"),
    ]
    for i, (col, head, body) in enumerate(items):
        y = 1.75 + i * 1.35
        shape(s, MSO_SHAPE.OVAL, 7.0, y + 0.08, 0.32, 0.32, fill=col)
        text(s, 7.55, y, 5.2, 0.45, head, size=20, bold=True)
        text(s, 7.55, y + 0.5, 5.2, 0.7, body, size=15, color=MUTED)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.0, 4.7, 5.75, 1.45, fill=CARD, name="Model")
    text(
        s,
        7.25,
        4.85,
        5.3,
        1.2,
        [
            [("comportamento", {"bold": True})],
            "= effetto rider + effetto infrastruttura",
            "+ interazione rider × infrastruttura",
        ],
        size=17,
        color=TEXT,
    )
    notes(
        s,
        "La struttura dei dati (una riga per tratto di percorso, con codice del rider e tipo di infrastruttura) "
        "permette di separare ciò che dipende dalla persona da ciò che dipende dal contesto. Eventuali stili di "
        "guida emergeranno dai dati, senza etichette a priori.",
    )


def slide_privacy(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Privacy & Etica")
    cards = [
        ("Nessuna manovra pericolosa", "nessuna frenata al limite, nessuna situazione provocata", icon_no),
        ("Guida naturalistica", "percorso reale, stile abituale, stop in qualsiasi momento", icon_route),
        ("Volti e targhe anonimizzati", "sfocatura automatica e controllo a campione", icon_blur),
        ("ID pseudonimi", "nessun nome nei dati; GPS di partenza e arrivo rimosso", icon_id),
    ]
    w, h, gx, gy, x0, y0 = 5.9, 1.85, 0.3, 0.3, 0.6, 1.6
    for i, (head, body, icon) in enumerate(cards):
        x = x0 + (i % 2) * (w + gx)
        y = y0 + (i // 2) * (h + gy)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=CARD, name=f"Card {head}")
        badge(s, x + 0.95, y + h / 2, 1.15)
        icon(s, x + 0.95, y + h / 2, 0.95)
        text(s, x + 1.8, y + 0.35, w - 2.0, 0.5, head, size=20, bold=True)
        text(s, x + 1.8, y + 0.9, w - 2.0, 0.8, body, size=15, color=MUTED)
    text(
        s,
        0.6,
        6.15,
        12.1,
        0.8,
        [
            [
                ("Documentazione per il Comitato Etico in preparazione: ", {"bold": True}),
                ("procedura, rischi, dati, persone riprese per caso. Le decisioni ancora aperte sono marcate TBD.", {}),
            ]
        ],
        size=15,
    )
    notes(
        s,
        "Sicurezza: nessuna situazione provocata, rispetto delle regole (casco obbligatorio). Privacy: codici al "
        "posto dei nomi, elaborazione locale, sfocatura di volti e targhe prima di guardare i video, protezione "
        "del GPS. Restano da definire con DPO e Comitato: base giuridica, DPIA, conservazione.",
    )


def slide_output(prs, overview_png: Path):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, ASPHALT)
    title(s, "Output", dark=True)
    items = [
        ("Metriche di comportamento", "velocità, decelerazioni, stop, stabilità"),
        ("Semantic gaze", "tempo di sguardo per categoria, scansione laterale"),
        ("Confronto tra infrastrutture", "stesso rider e rider diversi"),
        ("Dataset e metodologia", "pipeline aperta e documentata, riproducibile"),
    ]
    for i, (head, body) in enumerate(items):
        y = 1.65 + i * 1.12
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, y, 5.3, 0.95, fill=CARD_DARK, name=f"Output {head}")
        shape(s, MSO_SHAPE.OVAL, 0.85, y + 0.33, 0.3, 0.3, fill=AMBER if i == 1 else TEAL)
        text(s, 1.35, y + 0.1, 4.4, 0.42, head, size=18, bold=True, color=WHITE)
        text(s, 1.35, y + 0.52, 4.4, 0.38, body, size=14, color=SOFT)
    img_h, img_w = plt.imread(overview_png).shape[:2]
    pic_w = 6.55
    pic_h = pic_w * img_h / img_w
    s.shapes.add_picture(str(overview_png), Inches(6.2), Inches(1.65), Inches(pic_w), Inches(pic_h)).name = "MVP figure"
    text(
        s,
        6.2,
        1.65 + pic_h + 0.12,
        pic_w,
        0.4,
        "MVP già funzionante su dati sintetici (Windows e Linux)",
        size=13,
        italic=True,
        color=SOFT,
    )
    text(
        s,
        0.6,
        6.25,
        5.3,
        0.8,
        [
            [
                ("Prossimo passo: ", {"bold": True, "color": AMBER}),
                ("pilota con Aria Gen 2 dopo l'approvazione etica.", {"color": WHITE}),
            ]
        ],
        size=17,
    )
    notes(
        s,
        "Cosa produce il progetto: metriche quantitative per tratto di percorso e per interazione, tempi di "
        "sguardo per categoria, confronti tra infrastrutture e tra rider, una pipeline documentata. Oggi il "
        "sistema funziona end-to-end su dati sintetici; il primo passo reale è il pilota.",
    )


def build() -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    scene = render_scene_image(ASSETS / "scene_gaze_synthetic.png")
    overview = ROOT / "docs" / "figures" / "synthetic_overview.png"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = "AriaGenEScooter - overview"
    prs.core_properties.subject = "Naturalistic e-scooter riding study with Project Aria Gen 2"
    prs.core_properties.author = "AriaGenEScooter"
    slide_idea(prs)
    slide_recording(prs)
    slide_where(prs)
    slide_analysis(prs, scene)
    slide_perception_action(prs)
    slide_comparisons(prs)
    slide_privacy(prs)
    slide_output(prs, overview)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
